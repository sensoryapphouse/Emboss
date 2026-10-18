#!/usr/bin/env python3
"""Render a transcriber questionnaire (Emboss/docs/questions-*.md) to a Word file.

    python3 Emboss/scripts/questions-to-docx.py Emboss/docs/questions-round-2-KD.md [out.docx]

The markdown is the source of record; this keeps the Word file in the same shape as round 1
(`emboss-questions-for-transcribers.docx`, 22 Sep 2026): Title, Heading 1 for parts, Heading 2
for numbered questions, tick boxes as text, small "Ref:" lines, and real Word tables for the
tick-box tables. Understood markdown: `#`/`##`/`###` headings, fenced code blocks, `| a | b |`
tables (first row = header), `<sub>…</sub>`, `**bold**`, `` `code` ``, `---`.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm

INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`)')


def add_inline(par, text, size=None, italic=False):
    """Write `text` into paragraph `par`, honouring **bold** and `code` runs."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith('**') and piece.endswith('**'):
            run = par.add_run(piece[2:-2]); run.bold = True
        elif piece.startswith('`') and piece.endswith('`'):
            run = par.add_run(piece[1:-1]); run.font.name = 'Courier New'
        else:
            run = par.add_run(piece)
        if size: run.font.size = Pt(size)
        if italic: run.italic = True


def render(md_path: Path, out_path: Path):
    doc = Document()
    style = doc.styles['Normal']; style.font.name = 'Calibri'; style.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2); s.top_margin = s.bottom_margin = Cm(2)

    lines = md_path.read_text(encoding='utf-8').split('\n')
    i = 0
    para_buf = []

    def flush():
        nonlocal para_buf
        if para_buf:
            p = doc.add_paragraph(); add_inline(p, ' '.join(para_buf)); para_buf = []

    while i < len(lines):
        line = lines[i]
        if line.startswith('```'):
            flush(); i += 1; block = []
            while i < len(lines) and not lines[i].startswith('```'):
                block.append(lines[i]); i += 1
            p = doc.add_paragraph()
            run = p.add_run('\n'.join(block)); run.font.name = 'Courier New'; run.font.size = Pt(10)
            p.paragraph_format.left_indent = Cm(1)
            i += 1; continue
        if line.startswith('|'):
            flush(); rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'-*', c) for c in cells): rows.append(cells)
                i += 1
            ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncol); t.style = 'Table Grid'
            for r, row in enumerate(rows):
                for c in range(ncol):
                    cell = t.cell(r, c); cell.text = ''
                    add_inline(cell.paragraphs[0], row[c] if c < len(row) else '', size=10)
                    if r == 0:
                        for run in cell.paragraphs[0].runs: run.bold = True
            doc.add_paragraph(); continue
        if line.startswith('# '):
            flush(); doc.add_heading(line[2:].strip(), level=0)
        elif line.startswith('## '):
            flush(); doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith('### '):
            flush(); doc.add_heading(line[4:].strip(), level=2)
        elif line.strip() == '---':
            flush()
        elif line.startswith('<sub>'):
            flush(); p = doc.add_paragraph()
            add_inline(p, re.sub(r'</?sub>', '', line).strip(), size=8, italic=True)
        elif line.strip() == '':
            flush()
        elif '☐' in line or line.startswith('**') or line.rstrip().endswith('____'):
            flush(); p = doc.add_paragraph(); add_inline(p, line.strip())   # tick rows and lettered sub-questions keep their own line
        else:
            para_buf.append(line.strip())
        i += 1
    flush()
    doc.save(out_path)


if __name__ == '__main__':
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix('.docx')
    render(src, dst)
    print(f'wrote {dst}')
